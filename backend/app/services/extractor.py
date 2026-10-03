import fitz  # PyMuPDF
from docx import Document
from pathlib import Path

def extract_text_from_pdf(file_path: Path) -> str:
    """Extract text from a PDF file using PyMuPDF."""
    text_parts = []
    try:
        doc = fitz.open(str(file_path))
        for page in doc:
            page_text = page.get_text("text")
            if page_text:
                text_parts.append(page_text)
        doc.close()
    except Exception as e:
        raise ValueError(f"Failed to read PDF file: {str(e)}")

    extracted = "\n".join(text_parts).strip()
    if not extracted:
        raise ValueError("We couldn't extract readable text from this PDF. OCR support is not included in this MVP.")
    
    return extracted

def extract_text_from_docx(file_path: Path) -> str:
    """Extract text from a DOCX file using python-docx."""
    text_parts = []
    try:
        doc = Document(str(file_path))
        for para in doc.paragraphs:
            para_text = para.text.strip()
            if para_text:
                text_parts.append(para_text)
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells]
                row_text = " | ".join([c for c in row_cells if c])
                if row_text:
                    text_parts.append(row_text)
    except Exception as e:
        raise ValueError(f"Failed to read DOCX file: {str(e)}")

    extracted = "\n".join(text_parts).strip()
    if not extracted:
        raise ValueError("We couldn't extract readable text from this DOCX document.")
    
    return extracted

def extract_text_from_plain_text(file_path: Path) -> str:
    """Extract text from code files or text notes."""
    try:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            content = f.read().strip()
    except Exception as e:
        raise ValueError(f"Failed to read text file: {str(e)}")

    if not content:
        raise ValueError("The uploaded text file is empty.")

    return content

def extract_text(file_path: Path | str) -> str:
    """
    Main extraction interface.
    Determines file type and extracts readable text.
    """
    path = Path(file_path)
    ext = path.suffix.lower()

    if ext == ".pdf":
        return extract_text_from_pdf(path)
    elif ext == ".docx":
        return extract_text_from_docx(path)
    elif ext in {".c", ".cpp", ".h", ".hpp", ".py", ".java", ".js", ".ts", ".txt", ".md"}:
        return extract_text_from_plain_text(path)
    else:
        raise ValueError(f"This file type is not supported yet ({ext}).")
